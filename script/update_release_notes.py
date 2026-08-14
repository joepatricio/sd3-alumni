import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_release_notes():
    doc_path = r'docs\50\Najeal_SD3_ReleaseNotes_50Ver.docx'
    doc = docx.Document(doc_path)

    # New Features Summary Data
    new_features = [
        "Main Admin Dashboard with real-time operational telemetry, interactive pie charts, event calendar, and PNG report export",
        "Admin User Management tab with multi-status tabs, batch range filters, quick status actions (Suspend/Ban/Restore), and CSV export",
        "Admin Content Management System (CMS) for content moderation (Bulletins, Events, Donations)",
        "Admin Donation Dashboard Analytics & Metrics overview",
        "Edit Bulletin Modal with image attachment handling",
        "Edit Event Modal with category assignment and RSVP system"
    ]

    for_improvements = [
        "Server-Side REST API Querying & Sorting Engine (queryParser.js)",
        "Generic REST API Selective Column Retrieval (include query params)",
        "Enhanced Login Session Security, token validation, and role-based route protection",
        "Prisma v7 ORM integration and authentication API refactoring"
    ]

    fixes = [
        "Updated .gitignore to exclude Vite build outputs and Prisma generated client files",
        "Standardized UTC timestamp formatting for database records"
    ]

    # Comprehensive Test Cases Data (June - August 2026 Features)
    test_cases = [
        # Main Dashboard Test Cases
        (
            "1",
            "Main Dashboard - Real-time Operational Telemetry & Stat Refresh",
            "System telemetry metrics (Pending Users, Bulletins, Events, Total Contributions, Active Donors) load and refresh cleanly on click.",
            "",
            ""
        ),
        (
            "2",
            "Main Dashboard - Action-Required Moderation Banner Navigation",
            "Moderation alert banner displays exact count of pending items and quick-action buttons route directly to filtered moderation tabs.",
            "",
            ""
        ),
        (
            "3",
            "Main Dashboard - Interactive User Analytics & Demographics Pie Charts",
            "Recharts pie charts display user status distribution, generational cohorts, and profile privacy with toggleable legend highlights.",
            "",
            ""
        ),
        (
            "4",
            "Main Dashboard - Export Visual Analytics Report (PNG)",
            "Clicking 'Export Visual Report' captures dashboard telemetry and downloads a high-resolution PNG image report.",
            "",
            ""
        ),
        (
            "5",
            "Main Dashboard - Interactive Event Calendar & Day Filter",
            "Interactive calendar highlights dates with scheduled events; clicking a day filters and displays corresponding event details.",
            "",
            ""
        ),

        # User Management Tab Test Cases
        (
            "6",
            "User Management - Status Filter Tab Switching",
            "Clicking status tabs (All, Official, Regular, Pending, Suspended, Banned, Deactivated) filters table entries accurately and updates URL query parameter.",
            "",
            ""
        ),
        (
            "7",
            "User Management - Batch Year Range Filtering (Preset & Custom)",
            "Selecting batch range presets (Past 3, Past 5, Past 10, Past 30, or Custom Min/Max) filters users matching specified graduation years.",
            "",
            ""
        ),
        (
            "8",
            "User Management - Search by Name & Administrative Reason",
            "Search input filters user records dynamically by user name or administrative record description note.",
            "",
            ""
        ),
        (
            "9",
            "User Management - Table Column Sorting",
            "Clicking table headers (Name, Email, Batch, Granted/Expiry Date) toggles ascending/descending sorting with direction indicators.",
            "",
            ""
        ),
        (
            "10",
            "User Management - User Status Quick Actions (Suspend with Duration & Notes)",
            "Quick action buttons (Suspend, Ban, Restore, Deactivate) open edit modal with preset duration options (3 days, 1 week, 1 month) and log reason note.",
            "",
            ""
        ),
        (
            "11",
            "User Management - Export Users Dataset to CSV",
            "Clicking 'Export CSV' generates and downloads a users_export.csv file containing the filtered user table dataset.",
            "",
            ""
        ),
        (
            "12",
            "User Management - Create User Modal (Lazy-Loaded)",
            "Hovering or clicking 'Create User' triggers lazy loading of CreateUserModal, enabling admin account creation.",
            "",
            ""
        ),
        (
            "13",
            "User Management - Admin User Profile Preview Link",
            "Clicking the preview eye icon routes to /admin/preview/user/:id to inspect user profile in administrative preview mode.",
            "",
            ""
        ),

        # Admin CMS & Moderation Test Cases
        (
            "14",
            "Admin CMS - Moderate Pending Bulletin Post (Approve/Reject)",
            "Bulletin status updates to 'Approved' or 'Rejected' in database and public feed visibility updates accordingly.",
            "",
            ""
        ),
        (
            "15",
            "Admin CMS - Moderate Pending Event Post (Approve/Reject)",
            "Event status updates dynamically in Admin Content Table; only approved events display on public board.",
            "",
            ""
        ),
        (
            "16",
            "Admin CMS - Filter Content by Moderation Status",
            "Table dynamically filters and displays only items matching selected status (Pending, Approved, Rejected).",
            "",
            ""
        ),
        (
            "17",
            "Admin Analytics - View Donation Dashboard Metrics",
            "Summary metrics (Total Funds Raised, Total Donors, Target Progress) render accurately with current database totals.",
            "",
            ""
        ),

        # Security & Authentication Test Cases
        (
            "18",
            "Enhanced Auth Security - Protected Route Access Control",
            "Unauthenticated user attempting to access protected admin routes is blocked and redirected to Admin Login.",
            "",
            ""
        ),
        (
            "19",
            "Enhanced Auth Security - Expired Session Handling",
            "Expired session token revokes access cleanly with HTTP 401 and prompts user to re-authenticate.",
            "",
            ""
        ),
        (
            "20",
            "Enhanced Auth Security - Role-Based Dashboard Access",
            "Regular user attempting to log into Admin Portal with non-admin credentials receives clear authorization error.",
            "",
            ""
        ),

        # REST API Query Engine Test Cases
        (
            "21",
            "Server-Side REST Querying - Sort Alumni Directory Server-Side",
            "REST API parses sort parameters via queryParser.js and returns pre-sorted results matching requested order.",
            "",
            ""
        ),
        (
            "22",
            "Server-Side REST Querying - Sort Bulletins & Events Server-Side",
            "API executes server-side ordering query and updates UI list display accordingly.",
            "",
            ""
        ),
        (
            "23",
            "Server-Side REST POST - Create Bulletin with Image Attachment & Validation",
            "Image uploads to server storage, bulletin creates with 'Pending' status, and empty form submissions are blocked.",
            "",
            ""
        ),
    ]

    def set_cell_text(cell, text, font_size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(font_size)
        run.font.bold = bold
        return run

    # 1. Update Summary Table (Table 0)
    if len(doc.tables) > 0:
        table_summary = doc.tables[0]
        
        while len(table_summary.rows) > 1:
            table_summary._tbl.remove(table_summary.rows[-1]._tr)

        max_rows = max(len(new_features), len(for_improvements), len(fixes))
        for i in range(max_rows):
            row = table_summary.add_row()
            nf_text = f"- {new_features[i]}" if i < len(new_features) else ""
            imp_text = f"- {for_improvements[i]}" if i < len(for_improvements) else ""
            fix_text = f"- {fixes[i]}" if i < len(fixes) else ""

            set_cell_text(row.cells[0], nf_text, font_size=12)
            set_cell_text(row.cells[1], imp_text, font_size=12)
            set_cell_text(row.cells[2], fix_text, font_size=12)

    # 2. Update Test Cases Table (Table 1)
    if len(doc.tables) > 1:
        table_tests = doc.tables[1]

        while len(table_tests.rows) > 1:
            table_tests._tbl.remove(table_tests.rows[-1]._tr)

        for tc_no, tc_test, tc_expected, tc_pass, tc_action in test_cases:
            row = table_tests.add_row()
            
            set_cell_text(row.cells[0], tc_no, font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[1], tc_test, font_size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            set_cell_text(row.cells[2], tc_expected, font_size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            set_cell_text(row.cells[3], tc_pass, font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[4], tc_action, font_size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.save(doc_path)
    print(f"Successfully updated Release Notes document with {len(test_cases)} test cases!")

if __name__ == "__main__":
    update_release_notes()
